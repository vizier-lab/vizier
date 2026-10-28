"""Generates the golden files read by src/agents/tools/office/fixture_tests.rs.

They are written by other tools (python-docx, openpyxl, xlwt and hand-written OOXML), not by this
crate's own writers, and are not committed. Usage:

    pip install python-docx openpyxl xlwt
    python3 scripts/gen-office-fixtures.py [out_dir]

The default out_dir is src/agents/tools/office/fixtures.
"""
import os, sys, zipfile, datetime
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import openpyxl, xlwt

out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src", "agents", "tools", "office", "fixtures")
os.makedirs(out, exist_ok=True)

# ---- table.docx (python-docx) ----
d = Document()
d.add_heading("Quarterly Report", level=1)
p = d.add_paragraph("Plain, ")
p.add_run("bold").bold = True
p.add_run(" and ")
p.add_run("italic").italic = True
p.add_run(" text.")
d.add_paragraph("Top item", style="List Bullet")
d.add_paragraph("Nested item", style="List Bullet 2")
t = d.add_table(rows=3, cols=2)
for r, row in enumerate([["Name", "Note"], ["alpha", "a|b"], ["beta", "plain"]]):
    for c, v in enumerate(row):
        t.cell(r, c).text = v
p = d.add_paragraph("See ")
rid = d.part.relate_to("https://example.com/docs", RT.HYPERLINK, is_external=True)
h = OxmlElement("w:hyperlink")
h.set(qn("r:id"), rid)
r = OxmlElement("w:r")
tx = OxmlElement("w:t")
tx.text = "the docs"
r.append(tx)
h.append(r)
p._p.append(h)
d.save(f"{out}/table.docx")

# ---- mixed.xlsx (openpyxl) ----
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Types"
ws["A1"], ws["B1"] = "Name", "Value"
ws["A2"], ws["B2"] = "text", "hello"
ws["A3"], ws["B3"] = "integer", 12
ws["A4"], ws["B4"] = "float", 2.5
ws["A5"], ws["B5"] = "bool", True
ws["A6"] = "date"
ws["B6"] = datetime.datetime(2026, 11, 1)
ws["B6"].number_format = "yyyy-mm-dd"
ws["A7"] = "empty"
ws["A8"] = "error"
ws["B8"] = "#DIV/0!"
ws["B8"].data_type = "e"
ws2 = wb.create_sheet("Offset")
ws2["C5"] = "x"
ws2["D6"] = 7
wb.save(f"{out}/mixed.xlsx")

# ---- formulas.xlsx (hand-written OOXML: cached values) ----
def c(ref, v=None, f=None, t=None):
    a = f' t="{t}"' if t else ""
    return f'<c r="{ref}"{a}>' + (f"<f>{f}</f>" if f else "") + (f"<v>{v}</v>" if v is not None else "") + "</c>"

def istr(ref, s):
    return f'<c r="{ref}" t="inlineStr"><is><t>{s}</t></is></c>'

ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
data = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet {ns}><sheetData>'
    f'<row r="1">{istr("A1","Item")}{istr("B1","Amount")}</row>'
    f'<row r="2">{istr("A2","a")}{c("B2",17)}</row>'
    f'<row r="3">{istr("A3","b")}{c("B3",25)}</row>'
    f'<row r="4">{istr("A4","sum")}{c("B4",42,"SUM(B2:B3)")}</row>'
    f'<row r="5">{istr("A5","label")}{c("B5","big","IF(B4&gt;40,&quot;big&quot;,&quot;small&quot;)","str")}</row>'
    f'<row r="6">{istr("A6","double")}{c("B6",None,"B2*2")}</row>'
    f'<row r="7">{istr("A7","broken")}{c("B7","#DIV/0!","B2/0","e")}</row>'
    f'</sheetData></worksheet>'
)
summary = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet {ns}><sheetData>'
    f'<row r="1">{istr("A1","Total from Data")}{c("B1",42,"Data!B4")}</row>'
    f'</sheetData></worksheet>'
)
wbxml = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook {ns} '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
    '<sheet name="Data" sheetId="1" r:id="rId1"/><sheet name="Summary" sheetId="2" r:id="rId2"/>'
    '</sheets></workbook>'
)
rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
wbrels = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
    f'<Relationship Id="rId2" Type="{rel}/worksheet" Target="worksheets/sheet2.xml"/></Relationships>'
)
rootrels = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
)
ct = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
    '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
    '</Types>'
)
with zipfile.ZipFile(f"{out}/formulas.xlsx", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("[Content_Types].xml", ct)
    z.writestr("_rels/.rels", rootrels)
    z.writestr("xl/workbook.xml", wbxml)
    z.writestr("xl/_rels/workbook.xml.rels", wbrels)
    z.writestr("xl/worksheets/sheet1.xml", data)
    z.writestr("xl/worksheets/sheet2.xml", summary)

# ---- legacy.xls (xlwt) ----
x = xlwt.Workbook()
s = x.add_sheet("Old")
s.write(0, 0, "Item"); s.write(0, 1, "Qty")
s.write(1, 0, "pens"); s.write(1, 1, 12)
s.write(2, 0, "ink"); s.write(2, 1, 2.5)
s.write(3, 0, "ok"); s.write(3, 1, True)
x.save(f"{out}/legacy.xls")
